import unittest
from pathlib import Path

from card_engine.models import RecipeParsed
from card_engine.renderer import render_recipe_docx


class TestRenderer(unittest.TestCase):

    def setUp(self):
        self.out_dir = Path(__file__).parent.parent / "test_render_output"
        self.out_dir.mkdir(exist_ok=True)

    def tearDown(self):
        for f in self.out_dir.glob("*.docx"):
            f.unlink()
        if self.out_dir.exists():
            try:
                self.out_dir.rmdir()
            except OSError:
                pass

    def _make_recipe(self) -> RecipeParsed:
        return RecipeParsed(
            title="Test Pasta",
            ingredients_raw=["200g pasta", "100g cheese", "1 tbsp oil"],
            method_raw=["Cook pasta.", "Drain and mix with cheese.", "Drizzle with oil."],
            total_time="20 minutes",
            notes="Simple weeknight dinner.",
            source_file="test",
            content_hash="abc123",
        )

    def test_docx_created(self):
        recipe = self._make_recipe()
        out_path = str(self.out_dir / "test_pasta.docx")
        render_recipe_docx(recipe, metrics=None, research=[], output_path=out_path)
        self.assertTrue(Path(out_path).exists())
        self.assertGreater(Path(out_path).stat().st_size, 0)

    def test_docx_contains_title(self):
        from docx import Document
        recipe = self._make_recipe()
        out_path = str(self.out_dir / "test_title.docx")
        render_recipe_docx(recipe, metrics=None, research=[], output_path=out_path)
        doc = Document(out_path)
        texts = [p.text for p in doc.paragraphs]
        self.assertTrue(any("Test Pasta" in t for t in texts))

    def test_docx_has_paragraphs(self):
        from docx import Document
        recipe = self._make_recipe()
        out_path = str(self.out_dir / "test_paragraphs.docx")
        render_recipe_docx(recipe, metrics=None, research=[], output_path=out_path)
        doc = Document(out_path)
        self.assertGreater(len(doc.paragraphs), 0)
