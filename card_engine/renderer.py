from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import Optional

from card_engine.models import RecipeParsed
from card_engine.researcher import ResearchResult

logger = logging.getLogger(__name__)

try:
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    _DOCX_AVAILABLE = True
except ImportError:
    _DOCX_AVAILABLE = False
    logger.warning("python-docx not available; renderer will be limited")

# ---------------------------------------------------------------------------
# Quantity parsing helpers
# ---------------------------------------------------------------------------

_FRACTIONS = {"½": 0.5, "¼": 0.25, "¾": 0.75, "⅔": 0.667, "⅓": 0.333, "⅛": 0.125}

_UNIT_TO_G = {
    "kg": 1000.0, "g": 1.0,
    "ml": 1.0, "l": 1000.0, "litre": 1000.0, "liter": 1000.0,
    "tbsp": 15.0, "tablespoon": 15.0, "tablespoons": 15.0,
    "tsp": 5.0, "teaspoon": 5.0, "teaspoons": 5.0,
    "cup": 240.0, "cups": 240.0,
    "oz": 28.35, "lb": 453.6,
}

# Typical weight in grams for common whole items
_ITEM_WEIGHTS_G = {
    "egg": 55.0, "eggs": 55.0,
    "onion": 120.0, "onions": 120.0,
    "garlic": 5.0, "garlics": 5.0, "clove": 5.0, "cloves": 5.0,
    "carrot": 80.0, "carrots": 80.0,
    "celery": 50.0, "stick": 50.0, "sticks": 50.0,
    "pepper": 120.0, "peppers": 120.0,
    "chilli": 15.0, "chillies": 15.0,
    "tomato": 100.0, "tomatoes": 100.0,
    "shallot": 40.0, "shallots": 40.0,
    "lemon": 100.0, "lemons": 100.0,
    "lime": 60.0, "limes": 60.0,
    "potato": 100.0, "potatoes": 100.0,
}


def _parse_quantity_g(line: str) -> float:
    """Best-effort parse of leading quantity from an ingredient line → grams."""
    text = line.lstrip("•").strip()

    # Replace unicode fractions
    frac_val = 0.0
    for f, v in _FRACTIONS.items():
        if f in text:
            frac_val = v
            text = text.replace(f, "").strip()
            break

    # Match leading number optionally followed by a unit
    m = re.match(r"^(\d+(?:\.\d+)?)\s*([a-zA-Z]+)?", text)
    if not m:
        return 0.0

    num = float(m.group(1)) + frac_val
    unit_raw = (m.group(2) or "").lower().rstrip("s") if m.group(2) else ""

    # Try standard unit conversion
    unit_plural = (m.group(2) or "").lower()
    if unit_plural in _UNIT_TO_G:
        return num * _UNIT_TO_G[unit_plural]
    if unit_raw + "s" in _UNIT_TO_G:
        return num * _UNIT_TO_G[unit_raw + "s"]
    if unit_raw in _UNIT_TO_G:
        return num * _UNIT_TO_G[unit_raw]

    # Unit looks like an item name (e.g. "4 eggs", "2 onions")
    rest = text[m.end():].strip().lower()
    first_word_raw = unit_plural
    first_word = unit_raw
    for key in (first_word_raw, first_word, first_word + "s"):
        if key in _ITEM_WEIGHTS_G:
            return num * _ITEM_WEIGHTS_G[key]
    # Check rest of line for item keywords
    for key, weight in _ITEM_WEIGHTS_G.items():
        if re.search(r"\b" + re.escape(key) + r"\b", rest):
            return num * weight

    # Pure count with no recognised unit — assume each ~50 g
    return num * 50.0


def _ing_key(name: str) -> str:
    """Normalised key for matching ingredient lines to research results."""
    return re.sub(r"[^a-z0-9 ]", "", name.lower()).strip()


def _compute_metrics_from_research(
    ingredients_raw: list[str],
    research: list[ResearchResult],
    target_portions: int,
) -> dict:
    """
    Use ingredient quantities + research data to compute:
    - total_cost_low, total_cost_high (£)
    - total_kcal
    - cost_per_portion_low, cost_per_portion_high
    - kcal_per_portion
    Returns empty dict if no useful data.
    """
    # Build lookup: normalised ingredient name → ResearchResult
    lookup: dict[str, ResearchResult] = {}
    for r in research:
        lookup[_ing_key(r.ingredient_name)] = r

    total_cost_low = 0.0
    total_cost_high = 0.0
    total_kcal = 0.0
    any_cost = False
    any_kcal = False

    for ing_line in ingredients_raw:
        qty_g = _parse_quantity_g(ing_line)
        if qty_g <= 0:
            continue

        # Find matching research result
        ing_clean = _ing_key(ing_line)
        matched: Optional[ResearchResult] = None
        for key, res in lookup.items():
            if key in ing_clean or ing_clean in key:
                matched = res
                break
        if matched is None:
            continue

        if matched.cost_low_avg > 0 or matched.cost_high_avg > 0:
            # cost_low_avg / cost_high_avg are in £/g
            total_cost_low += matched.cost_low_avg * qty_g
            total_cost_high += matched.cost_high_avg * qty_g
            any_cost = True

        if matched.kcal_per_100g > 0:
            total_kcal += (matched.kcal_per_100g / 100.0) * qty_g
            any_kcal = True

    if not any_cost and not any_kcal:
        return {}

    portions = target_portions if target_portions > 0 else 1
    return {
        "total_cost_low": total_cost_low,
        "total_cost_high": total_cost_high,
        "cost_per_portion_low": total_cost_low / portions,
        "cost_per_portion_high": total_cost_high / portions,
        "total_kcal": total_kcal,
        "kcal_per_portion": total_kcal / portions,
        "any_cost": any_cost,
        "any_kcal": any_kcal,
    }


def _add_bold_heading(doc, text: str, level: int = 2) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(13 if level == 1 else 11)


def render_recipe_docx(
    recipe: RecipeParsed,
    metrics: Optional[dict],
    research: list[ResearchResult],
    output_path: str,
    target_portions: int = 10,
) -> None:
    """Create a .docx file for a single recipe."""
    if not _DOCX_AVAILABLE:
        raise RuntimeError("python-docx is not installed")

    doc = Document()

    # RECIPE CARD heading
    title_heading = doc.add_heading("RECIPE CARD", level=1)
    title_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Title
    _add_bold_heading(doc, "Title")
    doc.add_paragraph(recipe.title)

    # Yield
    _add_bold_heading(doc, "Yield")
    total_g = 0.0
    _weight_re = re.compile(r"(\d+(?:\.\d+)?)\s*(g|kg)\b", re.IGNORECASE)
    for ing in recipe.ingredients_raw:
        m = _weight_re.search(ing)
        if m:
            val = float(m.group(1))
            if m.group(2).lower() == "kg":
                val *= 1000
            total_g += val

    if total_g > 0:
        per_portion = total_g / target_portions
        yield_text = f"~{target_portions} portions | {total_g:.0f}g total | {per_portion:.0f}g per portion"
    else:
        yield_text = f"~{target_portions} portions | [total weight unknown] | [per portion unknown]"
    doc.add_paragraph(yield_text)

    # INGREDIENTS
    _add_bold_heading(doc, "INGREDIENTS")
    for ing in recipe.ingredients_raw:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(ing)

    # METHOD
    _add_bold_heading(doc, "METHOD")
    for i, step in enumerate(recipe.method_raw, 1):
        p = doc.add_paragraph(style="List Number")
        p.add_run(step)

    # TOTAL TIME
    _add_bold_heading(doc, "TOTAL TIME")
    doc.add_paragraph(recipe.total_time or "[unknown]")

    # Compute quantity-weighted metrics from research results
    computed = _compute_metrics_from_research(recipe.ingredients_raw, research, target_portions)

    # KCAL
    _add_bold_heading(doc, "KCAL")
    if metrics and metrics.get("kcal_per_serving", 0) > 0:
        kcal_text = f"{metrics['kcal_per_serving']:.0f} kcal per portion (estimated)"
    elif computed.get("any_kcal"):
        kcal_text = (
            f"~{computed['kcal_per_portion']:.0f} kcal per portion "
            f"({computed['total_kcal']:.0f} kcal total for {target_portions} portions, estimated)"
        )
    else:
        kcal_result = next((r for r in research if not r.todo and r.kcal_per_100g > 0), None)
        if kcal_result:
            kcal_text = f"~{kcal_result.kcal_per_100g:.0f} kcal per 100g (main ingredient, estimated)"
        else:
            todo_note = research[0].kcal_notes if research else "no data"
            kcal_text = f"TODO: [kcal could not be researched — {todo_note}]"
    doc.add_paragraph(kcal_text)

    # APPROXIMATE COST IN UK TODAY
    _add_bold_heading(doc, "APPROXIMATE COST IN UK TODAY")
    if metrics and metrics.get("total_cost", 0) > 0:
        total = metrics["total_cost"]
        cost_text = f"£{total:.2f} total | £{metrics['cost_per_serving']:.2f} per portion (estimated)"
    elif computed.get("any_cost"):
        lo = computed["total_cost_low"]
        hi = computed["total_cost_high"]
        plo = computed["cost_per_portion_low"]
        phi = computed["cost_per_portion_high"]
        cost_text = (
            f"£{lo:.2f}–£{hi:.2f} total | "
            f"£{plo:.2f}–£{phi:.2f} per portion "
            f"(estimated from UK supermarket data, {target_portions} portions)"
        )
    else:
        cost_results = [r for r in research if not r.todo and (r.cost_low_avg > 0 or r.cost_high_avg > 0)]
        if cost_results:
            low = sum(r.cost_low_avg for r in cost_results)
            high = sum(r.cost_high_avg for r in cost_results)
            cost_text = f"£{low:.2f}–£{high:.2f} (rate sum, estimated)"
        else:
            todo_note = research[0].cost_notes if research else "no data"
            cost_text = f"TODO: [cost could not be researched — {todo_note}]"
    doc.add_paragraph(cost_text)

    # FINISH
    _add_bold_heading(doc, "FINISH")
    finish_notes = ""
    m = re.search(r"\[FINISH:\s*([^\]]+)\]", recipe.notes or "")
    if m:
        finish_notes = m.group(1)
    doc.add_paragraph(finish_notes or "[see method]")

    # NOTES
    _add_bold_heading(doc, "NOTES")
    notes_parts = []
    if recipe.notes:
        clean_notes = re.sub(r"\[[A-Z ]+:\s*[^\]]+\]", "", recipe.notes).strip()
        if clean_notes:
            notes_parts.append(clean_notes)
    for r in research:
        if r.blocked_sources:
            notes_parts.append(f"Blocked sources for {r.ingredient_name}: {', '.join(r.blocked_sources)}")
        if r.todo:
            notes_parts.append(f"Research TODO: {r.ingredient_name} — {r.cost_notes}")
    doc.add_paragraph("\n".join(notes_parts) if notes_parts else "[none]")

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    logger.info("Saved recipe docx: %s", output_path)
