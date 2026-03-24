"""Export a RecipeCard to a .docx file (one file per recipe)."""
from __future__ import annotations

import os
from pathlib import Path

from docx import Document
from docx.shared import Pt

from card_engine.models import RecipeCard


def export_docx(card: RecipeCard, out_dir: str | os.PathLike) -> Path:
    """
    Write *card* to ``<out_dir>/<slug>.docx`` and return the output path.

    Layout (metric-only):
      1. Title  (Heading 1)
      2. Yield / Portions
      3. Cost summary  (low / mid / high)
      4. Kcal summary  (low / mid / high)
      5. Ingredients table
      6. Method steps
      7. Notes (if any)
    """
    out_path = Path(out_dir) / f"{card.slug}.docx"
    out_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()

    # 1. Title
    doc.add_heading(card.title, level=1)

    # 2. Yield / Portions
    yield_text = f"{card.portions} portions"
    if card.total_yield_g:
        yield_text += f"  |  Total yield: {card.total_yield_g:.0f} g/ml"
    doc.add_paragraph(yield_text)

    # 3. Cost summary
    doc.add_heading("Cost", level=2)
    cost_para = doc.add_paragraph()
    cost_para.add_run(
        f"Low: £{card.cost_low:.2f}  |  Mid: £{card.cost_mid:.2f}"
        f"  |  High: £{card.cost_high:.2f}\n"
    )
    cost_para.add_run(f"Per portion (mid): £{card.cost_per_portion_mid:.2f}")

    # 4. Kcal summary
    doc.add_heading("Calories", level=2)
    kcal_para = doc.add_paragraph()
    kcal_para.add_run(
        f"Low: {card.kcal_low:.0f} kcal  |  Mid: {card.kcal_mid:.0f} kcal"
        f"  |  High: {card.kcal_high:.0f} kcal\n"
    )
    kcal_para.add_run(f"Per portion (mid): {card.kcal_per_portion_mid:.0f} kcal")

    # 5. Ingredients table
    doc.add_heading("Ingredients", level=2)
    if card.ingredients:
        tbl = doc.add_table(rows=1, cols=3)
        tbl.style = "Table Grid"
        hdr = tbl.rows[0].cells
        hdr[0].text = "Ingredient"
        hdr[1].text = "Quantity"
        hdr[2].text = "Unit"
        for ing in card.ingredients:
            row = tbl.add_row().cells
            row[0].text = ing.name
            row[1].text = str(ing.quantity) if ing.quantity is not None else ""
            row[2].text = ing.unit or ""
    else:
        doc.add_paragraph("No ingredients listed.")

    # 6. Method
    doc.add_heading("Method", level=2)
    for i, step in enumerate(card.method, start=1):
        doc.add_paragraph(f"{i}. {step}")

    # 7. Notes
    if card.notes:
        doc.add_heading("Notes", level=2)
        for note in card.notes:
            doc.add_paragraph(f"• {note}")

    doc.save(str(out_path))
    return out_path
