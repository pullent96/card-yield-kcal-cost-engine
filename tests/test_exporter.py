"""Tests for DOCX exporter."""
from __future__ import annotations

import tempfile
from pathlib import Path

import pytest
from docx import Document

from card_engine.exporter import export_docx
from card_engine.models import (
    CostResult,
    IngredientLine,
    KcalResult,
    Range,
    RecipeCard,
    YieldResult,
)


def _make_card(title="Test Recipe", slug="test-recipe") -> RecipeCard:
    return RecipeCard(
        title=title,
        slug=slug,
        portions=4,
        total_yield_g=800.0,
        cost_low=2.00,
        cost_mid=3.00,
        cost_high=4.00,
        cost_per_portion_mid=0.75,
        kcal_low=400.0,
        kcal_mid=600.0,
        kcal_high=800.0,
        kcal_per_portion_mid=150.0,
        ingredients=[
            IngredientLine("flour", "500g flour", 500.0, "g", "flour"),
            IngredientLine("water", "300ml water", 300.0, "ml", "water"),
        ],
        method=["Mix flour and water.", "Bake at 180°C for 30 minutes."],
        notes=["Best served warm."],
        cost_breakdown=[
            CostResult(
                ingredient_key="flour",
                used_cost_mid=1.50,
                range=Range(low=1.20, mid=1.50, high=1.80),
            )
        ],
        kcal_breakdown=[
            KcalResult(
                ingredient_key="flour",
                total_kcal_mid=400.0,
                range=Range(low=380.0, mid=400.0, high=420.0),
            )
        ],
    )


class TestExportDocx:
    def test_file_created(self):
        with tempfile.TemporaryDirectory() as tmp:
            card = _make_card()
            path = export_docx(card, tmp)
            assert path.exists()
            assert path.suffix == ".docx"

    def test_filename_uses_slug(self):
        with tempfile.TemporaryDirectory() as tmp:
            card = _make_card(title="Chilli Con Carne", slug="chilli-con-carne")
            path = export_docx(card, tmp)
            assert path.name == "chilli-con-carne.docx"

    def test_title_in_document(self):
        with tempfile.TemporaryDirectory() as tmp:
            card = _make_card(title="My Special Recipe")
            path = export_docx(card, tmp)
            doc = Document(str(path))
            headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
            assert any("My Special Recipe" in h for h in headings)

    def test_ingredients_present(self):
        with tempfile.TemporaryDirectory() as tmp:
            card = _make_card()
            path = export_docx(card, tmp)
            doc = Document(str(path))
            # Ingredients are in a table; check table cell text
            all_text = " ".join(
                cell.text for table in doc.tables for row in table.rows for cell in row.cells
            )
            assert "flour" in all_text.lower()

    def test_method_steps_present(self):
        with tempfile.TemporaryDirectory() as tmp:
            card = _make_card()
            path = export_docx(card, tmp)
            doc = Document(str(path))
            full_text = " ".join(p.text for p in doc.paragraphs)
            assert "Mix flour and water" in full_text

    def test_notes_present(self):
        with tempfile.TemporaryDirectory() as tmp:
            card = _make_card()
            path = export_docx(card, tmp)
            doc = Document(str(path))
            full_text = " ".join(p.text for p in doc.paragraphs)
            assert "Best served warm" in full_text

    def test_no_notes_when_empty(self):
        with tempfile.TemporaryDirectory() as tmp:
            card = _make_card()
            card.notes = []
            path = export_docx(card, tmp)
            doc = Document(str(path))
            headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
            assert not any("Notes" in h for h in headings)

    def test_one_file_per_recipe(self):
        """Exporting two recipes creates two separate .docx files."""
        with tempfile.TemporaryDirectory() as tmp:
            card1 = _make_card(title="Recipe A", slug="recipe-a")
            card2 = _make_card(title="Recipe B", slug="recipe-b")
            path1 = export_docx(card1, tmp)
            path2 = export_docx(card2, tmp)
            assert path1 != path2
            assert path1.exists()
            assert path2.exists()

    def test_output_dir_created_if_missing(self):
        with tempfile.TemporaryDirectory() as tmp:
            new_dir = Path(tmp) / "new_subdir"
            card = _make_card()
            path = export_docx(card, new_dir)
            assert path.exists()
